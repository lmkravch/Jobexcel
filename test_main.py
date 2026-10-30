import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import generate_examples
from main import COLUMNS, normalize_phone, parse_price, process


class ExcelTests(unittest.TestCase):
    def test_complete_workflow(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch.object(generate_examples, 'BASE_DIR', root):
                generate_examples.generate()
                with self.assertRaises(SystemExit):
                    generate_examples.generate()
            self.assertEqual(len(list((root / 'input').glob('*.xlsx'))), 5)
            output = root / 'result.xlsx'
            self.assertEqual(process(root / 'input', output), (16, 3, 8, 5))
            self.assertTrue(output.is_file())
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ['Данные', 'Ошибки', 'Итоги'])
            clean = list(book['Данные'].values)[1:]
            errors = list(book['Ошибки'].values)[1:]
            self.assertEqual(len(clean), len(set(clean)))
            self.assertTrue(all(re.fullmatch(r'\+7\d{10}', r[2]) for r in clean))
            self.assertTrue(all(r[2] and r[4] is not None for r in clean))
            self.assertEqual(sum(r[2] is None for r in errors), 3)
            self.assertEqual(sum(r[4] is None for r in errors), 3)
            self.assertTrue(all(r[5] and r[6] and r[7] >= 2 for r in errors))
            self.assertEqual(dict(list(book['Итоги'].values)[1:]),
                             {'Анна': 600.6, 'Борис': 500, 'Вера': 75.5, 'Дарья': 25})
            book.close()
            self.assertEqual(process(root / 'input', output), (16, 3, 8, 5))

    def test_phone_and_price_validation(self):
        for phone in ['+7 999 123-45-67', '89991234567', '9991234567', 89991234567.0]:
            self.assertEqual(normalize_phone(phone), '+79991234567')
        for phone in [None, '', '123', '+19991234567', 'abc9991234567']:
            with self.assertRaises(ValueError):
                normalize_phone(phone)
        self.assertEqual(str(parse_price('100,50')), '100.50')
        self.assertEqual(parse_price(0), 0)
        for price in [None, '', 'не число', 'NaN', 'Infinity', -1, '1.001']:
            with self.assertRaises(ValueError):
                parse_price(price)

    def test_missing_files_and_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with self.assertRaisesRegex(ValueError, 'нет файлов'):
                process(root, root / 'out.xlsx')
            book = Workbook()
            book.active.append(['Клиент'])
            book.save(root / 'bad.xlsx')
            book.close()
            with self.assertRaisesRegex(ValueError, 'нужны колонки'):
                process(root, root / 'out.xlsx')


if __name__ == '__main__':
    unittest.main()
