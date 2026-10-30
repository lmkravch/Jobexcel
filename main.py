"""Объединение Excel-файлов, проверка данных и подсчёт покупок."""
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
COLUMNS = ['Дата', 'Клиент', 'Телефон', 'Товар', 'Цена']


def normalize_phone(value):
    if value is None or not str(value).strip():
        raise ValueError('Нет телефона')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not re.fullmatch(r'\+?[0-9\s()\-]+', text):
        raise ValueError('Некорректный телефон')
    digits = re.sub(r'\D', '', text)
    if len(digits) == 10:
        digits = '7' + digits
    elif len(digits) == 11 and digits[0] == '8':
        digits = '7' + digits[1:]
    if len(digits) != 11 or not digits.startswith('7'):
        raise ValueError('Некорректный телефон')
    return '+' + digits


def parse_price(value):
    if value is None or not str(value).strip():
        raise ValueError('Нет цены')
    try:
        price = Decimal(str(value).replace(' ', '').replace(',', '.'))
    except InvalidOperation:
        raise ValueError('Некорректная цена') from None
    if not price.is_finite() or price < 0 or price != price.quantize(Decimal('0.01')):
        raise ValueError('Некорректная цена')
    return price


def process(input_dir=BASE_DIR / 'input', output=BASE_DIR / 'result.xlsx'):
    files = sorted(Path(input_dir).glob('*.xlsx'))
    files = [p for p in files if not p.name.startswith('~$')]
    if not files:
        raise ValueError('В папке input нет файлов .xlsx')
    clean, errors, seen = [], [], set()
    totals = defaultdict(Decimal)
    count = duplicates = 0
    for path in files:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = book.active.iter_rows(values_only=True)
            headers = next(rows, ())
            if any(col not in headers for col in COLUMNS):
                raise ValueError(f'{path.name}: нужны колонки {", ".join(COLUMNS)}')
            indices = [headers.index(col) for col in COLUMNS]
            for number, source in enumerate(rows, 2):
                row = tuple(source[i] if i < len(source) else None for i in indices)
                if all(v is None for v in row):
                    continue
                count += 1
                # Полные дубли сравниваем до преобразования телефонов.
                if row in seen:
                    duplicates += 1
                    continue
                seen.add(row)
                problems = []
                try:
                    phone = normalize_phone(row[2])
                except ValueError as exc:
                    problems.append(str(exc))
                try:
                    price = parse_price(row[4])
                except ValueError as exc:
                    problems.append(str(exc))
                if row[1] is None or not str(row[1]).strip():
                    problems.append('Нет клиента')
                if problems:
                    errors.append(list(row) + ['; '.join(problems), path.name, number])
                    continue
                clean.append([row[0], row[1], phone, row[3], float(price)])
                totals[str(row[1])] += price
        finally:
            book.close()
    result = Workbook()
    result.remove(result.active)
    for name, headers, rows in [
        ('Данные', COLUMNS, clean),
        ('Ошибки', COLUMNS + ['Причина', 'Файл', 'Строка'], errors),
        ('Итоги', ['Клиент', 'Общая сумма покупок'],
         [(client, float(total)) for client, total in sorted(totals.items())]),
    ]:
        sheet = result.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
            # Любой исходный текст сохраняем как текст, включая строки с '='.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    result.save(output)
    result.close()
    print(f'Прочитано: {count}; удалено дублей: {duplicates}; '
          f'данных: {len(clean)}; ошибок: {len(errors)}. Создан {output}')
    return count, duplicates, len(clean), len(errors)


if __name__ == '__main__':
    try:
        process()
    except (ValueError, OSError) as exc:
        raise SystemExit(f'Ошибка: {exc}') from exc
